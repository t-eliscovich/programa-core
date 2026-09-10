"""Sus datos: lo que el cliente carga en el primer ingreso al portal (mig
0249, dueña 10/09/2026): correo para facturas, celular, fijo y la dirección
de entrega SEGMENTADA. Obligatorio; "que no lo dejen vacío o alguna
restricción así". Y la pantalla de la oficina que lo mira, lo pasa a la
ficha (correo y celular) y lo exporta a Excel para Asinfo.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from modules.portal import datos  # noqa: E402

FICHA = {"codigo_cli": "AJT", "nombre": "TOTOY BUITRON ANDRES JULIO", "ruc": "1724354004001",
         "vend": "EDG", "correo": "", "telefono": "0997857539",
         "direccion1": "SOLANO DE QUIÑONEZ OE1-126 Y ALFARO FTE PARQUE ECOLOGICO",
         "direccion2": "FARO FTE PARQUE ECOLOGICO", "provincia": "PICHINCHA", "canton": "QUITO"}

BUENO = {"correo_facturas": "Facturas@Totoy.com", "celular": "099 785 7539", "telefono": "",
         "ciudad": "Quito", "calle_principal": "Solano de Quiñonez", "numero": "OE1-126",
         "calle_secundaria": "Alfaro", "referencia": "Frente al parque"}


def _app_portal():
    from tests.test_routes_smoke import build_app
    with patch.dict(os.environ, {**os.environ, "MODO": "portal"}):
        app, deshacer = build_app()
    app.config["WTF_CSRF_ENABLED"] = False
    return app, deshacer


def _sesion(app):
    c = app.test_client()
    with c.session_transaction() as s:
        s["portal_cliente"] = "AJT"
    return c


def _ficha(monkeypatch):
    from modules.portal import acceso
    monkeypatch.setattr(acceso, "ficha", lambda cod: dict(FICHA))
    monkeypatch.setattr(acceso, "cliente", lambda cod: dict(FICHA))
    monkeypatch.setattr(acceso, "acceso", lambda cod: {"mail": "teliscovich@gmail.com", "activo": True, "clave_hash": "x"})


# ---------------------------------------------------------------------------
# Las reglas
# ---------------------------------------------------------------------------


def test_lo_bueno_pasa_y_sale_normalizado():
    d, errores = datos.validar(BUENO)
    assert errores == {}
    assert d["celular"] == "0997857539" and d["correo_facturas"] == "facturas@totoy.com"
    assert d["telefono"] == "" and d["numero"] == "OE1-126"


def test_vacio_no_pasa_y_dice_que_campo():
    d, errores = datos.validar({})
    assert set(errores) == set(datos.OBLIGATORIOS)
    assert all(v == "Falta completar." for v in errores.values())


def test_el_correo_el_celular_y_el_fijo_tienen_forma():
    _, e = datos.validar({**BUENO, "correo_facturas": "totoy.com"})
    assert e == {"correo_facturas": "No parece un correo."}
    _, e = datos.validar({**BUENO, "celular": "022222222"})
    assert "celular" in e and "09" in e["celular"]
    _, e = datos.validar({**BUENO, "celular": "+593 99 785 7539"})
    assert e == {}
    _, e = datos.validar({**BUENO, "telefono": "12"})
    assert "telefono" in e
    d, e = datos.validar({**BUENO, "telefono": "(02) 245-6789"})
    assert e == {} and d["telefono"] == "022456789"
    # "S/N" vale como número.
    _, e = datos.validar({**BUENO, "numero": "S/N"})
    assert e == {}


def test_la_precarga_trae_lo_que_teniamos_sin_adivinar_la_direccion():
    d = datos.precarga(FICHA, "teliscovich@gmail.com")
    assert d["correo_facturas"] == "teliscovich@gmail.com"
    assert d["celular"] == "0997857539" and d["telefono"] == ""
    assert d["ciudad"] == "Quito"
    # La dirección vieja va entera en calle principal; número y calle B vacíos.
    assert d["calle_principal"] == FICHA["direccion1"] and d["numero"] == "" and d["calle_secundaria"] == ""
    # direccion2 que repite la cola de direccion1 no se precarga como referencia.
    assert d["referencia"] == ""
    # Un fijo en la ficha va al fijo, no al celular.
    d2 = datos.precarga({**FICHA, "telefono": "022456789"})
    assert d2["celular"] == "" and d2["telefono"] == "022456789"


def test_la_direccion_en_una_linea_se_escribe_como_en_ecuador():
    assert datos.direccion_en_una_linea({"calle_principal": "Solano", "numero": "OE1-126",
                                         "calle_secundaria": "Alfaro", "referencia": "Frente al parque"}) \
        == "Solano OE1-126 y Alfaro, Frente al parque"


# ---------------------------------------------------------------------------
# El freno y la pantalla
# ---------------------------------------------------------------------------


def test_sin_sus_datos_todo_manda_a_la_pantalla_y_salir_sigue_andando(monkeypatch):
    app, deshacer = _app_portal()
    try:
        _ficha(monkeypatch)
        monkeypatch.setattr(datos, "completo", lambda cod: False)
        c = _sesion(app)
        for ruta in ("/", "/estado-de-cuenta", "/facturas", "/mis-pagos", "/mis-datos"):
            r = c.get(ruta)
            assert r.status_code == 302 and r.headers["Location"].endswith("/sus-datos"), ruta
        assert c.get("/sus-datos").status_code == 200
        assert c.get("/salir").status_code in (200, 302)
    finally:
        deshacer()


def test_la_pantalla_viene_precargada_y_la_primera_vez_es_obligatoria(monkeypatch):
    app, deshacer = _app_portal()
    try:
        _ficha(monkeypatch)
        monkeypatch.setattr(datos, "completo", lambda cod: False)
        monkeypatch.setattr(datos, "leer", lambda cod: None)
        html = _sesion(app).get("/sus-datos").get_data(as_text=True)
        assert "Antes de entrar, sus datos" in html and "Guardar y entrar" in html
        assert 'value="0997857539"' in html and 'value="Quito"' in html
        assert 'value="teliscovich@gmail.com"' in html
        assert "‹ Perfil" not in html  # sin salida: es obligatorio
    finally:
        deshacer()


def test_mandar_vacio_vuelve_con_los_campos_marcados_y_no_guarda(monkeypatch):
    app, deshacer = _app_portal()
    try:
        _ficha(monkeypatch)
        monkeypatch.setattr(datos, "leer", lambda cod: None)
        guardados = []
        monkeypatch.setattr(datos, "guardar", lambda cod, d, previo=None: guardados.append(d))
        r = _sesion(app).post("/sus-datos", data={"correo_facturas": "x", "celular": "123"})
        assert r.status_code == 400
        html = r.get_data(as_text=True)
        assert "Revise los campos marcados" in html
        assert "No parece un correo." in html and "Falta completar." in html
        assert guardados == []
    finally:
        deshacer()


def test_con_todo_bien_guarda_con_la_foto_de_la_ficha_y_entra(monkeypatch):
    app, deshacer = _app_portal()
    try:
        _ficha(monkeypatch)
        monkeypatch.setattr(datos, "leer", lambda cod: None)
        guardados = []
        monkeypatch.setattr(datos, "guardar", lambda cod, d, previo=None: guardados.append((cod, d, previo)))
        r = _sesion(app).post("/sus-datos", data=BUENO)
        assert r.status_code == 302 and r.headers["Location"].endswith("/")
        cod, d, previo = guardados[0]
        assert cod == "AJT" and d["celular"] == "0997857539" and previo["direccion1"] == FICHA["direccion1"]
    finally:
        deshacer()


def test_despues_se_edita_desde_perfil_y_perfil_muestra_lo_cargado(monkeypatch):
    app, deshacer = _app_portal()
    try:
        _ficha(monkeypatch)
        cargado = {**BUENO, "celular": "0997857539", "correo_facturas": "facturas@totoy.com"}
        monkeypatch.setattr(datos, "leer", lambda cod: dict(cargado))
        c = _sesion(app)
        perfil = c.get("/mis-datos").get_data(as_text=True)
        assert "facturas@totoy.com" in perfil and "Solano de Quiñonez OE1-126 y Alfaro" in perfil
        assert 'href="/sus-datos"' in perfil
        editar = c.get("/sus-datos").get_data(as_text=True)
        assert "‹ Perfil" in editar and 'value="OE1-126"' in editar and ">Guardar<" in editar
    finally:
        deshacer()


def test_despues_de_elegir_la_clave_va_a_sus_datos():
    fuente = (ROOT / "modules" / "portal" / "views.py").read_text(encoding="utf-8")
    bloque = fuente.split("def elegir_clave(")[1].split("\n@portal_bp.route")[0]
    assert 'redirect(url_for("portal.sus_datos"))' in bloque


def test_si_la_base_no_contesta_el_freno_deja_pasar(monkeypatch):
    """Nunca todos los clientes afuera por plomería (la mig 0249 sin correr)."""
    import db
    monkeypatch.setattr(db, "fetch_one", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no table")))
    monkeypatch.undo()  # el conftest pisa `completo`; acá se prueba el de verdad
    monkeypatch.setattr(db, "fetch_one", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no table")))
    assert datos.completo("AJT") is True


# ---------------------------------------------------------------------------
# La oficina
# ---------------------------------------------------------------------------


FILA = {**BUENO, "codigo_cli": "AJT", "nombre": "TOTOY", "vend": "EDG", "celular": "0997857539",
        "correo_facturas": "facturas@totoy.com", "ficha_correo": "", "ficha_telefono": "099-7857539",
        "ficha_direccion1": FICHA["direccion1"], "ficha_direccion2": "", "ficha_canton": "QUITO",
        "actualizado_en": None, "aplicado_en": None, "aplicado_por": None}


def test_las_diferencias_se_marcan_por_campo():
    dif = datos.con_diferencias(FILA)
    assert dif == {"correo": True, "telefono": False, "direccion": True}


def test_pasar_a_la_ficha_solo_correo_y_celular(monkeypatch):
    from modules.clientes import queries as cq
    llamadas, marcados = [], []
    monkeypatch.setattr(datos, "leer", lambda cod: dict(FILA))
    monkeypatch.setattr(cq, "editar", lambda cod, **kw: llamadas.append((cod, kw)) or 1)
    import db
    monkeypatch.setattr(db, "execute", lambda sql, params=None, conn=None: marcados.append(params))
    assert datos.pasar_a_la_ficha("AJT", "tamara") == 1
    cod, kw = llamadas[0]
    assert cod == "AJT" and kw == {"correo": "facturas@totoy.com", "telefono": "0997857539", "usuario": "tamara"}
    assert "direccion1" not in kw  # la dirección va por Asinfo: el sync la pisa
    assert marcados[0] == ("tamara", "AJT")


def test_el_excel_va_segmentado_para_asinfo():
    from io import BytesIO

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(datos.excel([FILA])))
    ws = wb.active
    assert [c.value for c in ws[1]] == [t for _, t in datos.COLUMNAS_EXCEL]
    fila = [c.value for c in ws[2]]
    assert fila[:10] == ["AJT", "TOTOY", "facturas@totoy.com", "0997857539", None, "Quito",
                         "Solano de Quiñonez", "OE1-126", "Alfaro", "Frente al parque"]


def test_la_pantalla_de_la_oficina_lista_y_pide_permiso(app, monkeypatch):
    monkeypatch.setattr(datos, "cargados", lambda: [dict(FILA)])

    @app.before_request
    def _entrar():  # pragma: no cover - infra de test
        from flask import g, session
        session["user_id"] = 1
        g.user = {"id_usuario": 1, "username": "tamara", "nombre_rol": "Accionista", "activo": True, "vend": None}
        g.permisos = {"clientes.ver", "clientes.editar"}

    html = app.test_client().get("/clientes/datos-del-portal").get_data(as_text=True)
    assert "AJT" in html and "facturas@totoy.com" in html and "Solano de Quiñonez OE1-126 y Alfaro" in html
    assert "Pasar correo y celular a la ficha" in html and "Excel para Asinfo" in html
    r = app.test_client().get("/clientes/datos-del-portal.xlsx")
    assert r.status_code == 200 and r.mimetype.endswith("sheet")
